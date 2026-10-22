"""Extract text data and rendered pixels from spreadsheets and vector images.

The source file remains one object.  This service only produces representations
of that object: table text, chart/vector text, and PNG renderings for the normal
OCR and image-semantic pipeline.
"""

from __future__ import annotations

import argparse
import gzip
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET


SPREADSHEET_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
VECTOR_EXTENSIONS = {".svg", ".svgz"}
_SAFE_NAME = re.compile(r"[^A-Za-z0-9_.-]+")


def _issue(stage: str, message: str, severity: str = "error") -> dict[str, str]:
    return {"stage": stage, "message": message, "severity": severity}


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value).replace("\r\n", "\n").replace("\r", "\n")


def _quoted(value: object) -> str:
    return json.dumps(_cell_text(value), ensure_ascii=False)


def _worksheet_table_text(sheet: object, cached_sheet: object | None) -> str | None:
    from openpyxl.utils import get_column_letter

    populated: list[tuple[int, list[tuple[int, str]]]] = []
    min_column: int | None = None
    max_column: int | None = None
    min_row: int | None = None
    max_row: int | None = None

    for row in sheet.iter_rows():
        values: list[tuple[int, str]] = []
        for cell in row:
            value = cell.value
            if value is None:
                continue
            text = _cell_text(value)
            if getattr(cell, "data_type", None) == "f" and cached_sheet is not None:
                cached = cached_sheet[cell.coordinate].value
                if cached is not None:
                    text = f"{text} => {_cell_text(cached)}"
            values.append((cell.column, text))
            min_column = cell.column if min_column is None else min(min_column, cell.column)
            max_column = cell.column if max_column is None else max(max_column, cell.column)
        if values:
            populated.append((row[0].row, values))
            min_row = row[0].row if min_row is None else min(min_row, row[0].row)
            max_row = row[0].row if max_row is None else max(max_row, row[0].row)

    if not populated or min_column is None or max_column is None:
        return None

    range_label = (
        f"{get_column_letter(min_column)}{min_row}:"
        f"{get_column_letter(max_column)}{max_row}"
    )
    lines = ["TABLE", f"sheet: {_quoted(sheet.title)}", f"used_range: {range_label}"]
    merged = [str(item) for item in sheet.merged_cells.ranges]
    if merged:
        lines.append("merged_ranges: " + ", ".join(merged))
    table_definitions = [
        f"{name}={reference}" for name, reference in sheet.tables.items()
    ]
    if table_definitions:
        lines.append("named_tables: " + ", ".join(table_definitions))
    lines.append("cells:")
    for row_index, values in populated:
        encoded = "\t".join(
            f"{get_column_letter(column)}{row_index}={_quoted(value)}"
            for column, value in values
        )
        lines.append(encoded)
    return "\n".join(lines)


def _nested(value: object, *names: str) -> object | None:
    current: object | None = value
    for name in names:
        if current is None:
            return None
        current = getattr(current, name, None)
    return current


def _first_nested(value: object, paths: tuple[tuple[str, ...], ...]) -> object | None:
    for path in paths:
        found = _nested(value, *path)
        if found not in (None, ""):
            return found
    return None


def _rich_text(value: object) -> str | None:
    fragments: list[str] = []
    for paragraph in _nested(value, "tx", "rich", "p") or []:
        for run in getattr(paragraph, "r", []) or []:
            text = getattr(run, "t", None)
            if text:
                fragments.append(str(text))
        plain = getattr(paragraph, "endParaRPr", None)
        if not fragments and plain is not None:
            text = getattr(paragraph, "t", None)
            if text:
                fragments.append(str(text))
    text = "".join(fragments).strip()
    return text or None


def _range_values(workbook: object, formula: object | None) -> list[object]:
    if not formula:
        return []
    from openpyxl.utils.cell import range_to_tuple

    expression = str(formula).lstrip("=")
    try:
        sheet_name, boundaries = range_to_tuple(expression)
        minimum_column, minimum_row, maximum_column, maximum_row = boundaries
        sheet = workbook[sheet_name]
    except (KeyError, TypeError, ValueError):
        return []
    values: list[object] = []
    for row in sheet.iter_rows(
        min_row=minimum_row,
        max_row=maximum_row,
        min_col=minimum_column,
        max_col=maximum_column,
    ):
        values.extend(cell.value for cell in row)
    return values


def _literal_cache(reference: object, cache_name: str) -> list[object]:
    cache = getattr(reference, cache_name, None)
    points = getattr(cache, "pt", None) if cache is not None else None
    if not points:
        return []
    ordered = sorted(points, key=lambda point: int(getattr(point, "idx", 0)))
    return [getattr(point, "v", None) for point in ordered]


def _reference_values(workbook: object, reference: object | None) -> tuple[str | None, list[object]]:
    if reference is None:
        return None, []
    formula = _first_nested(
        reference,
        (("numRef", "f"), ("strRef", "f")),
    )
    values = _range_values(workbook, formula)
    if not values:
        values = _literal_cache(reference, "numLit") or _literal_cache(reference, "strLit")
    return (str(formula) if formula else None), values


def _series_name(workbook: object, series: object, index: int) -> str:
    title = getattr(series, "tx", None)
    literal = getattr(title, "v", None) if title is not None else None
    if literal:
        return str(literal)
    formula = _nested(title, "strRef", "f")
    values = _range_values(workbook, formula)
    if values and values[0] is not None:
        return str(values[0])
    return f"Series {index}"


def _chart_payload(workbook: object, sheet: object, chart: object, index: int) -> dict[str, object]:
    chart_type = type(chart).__name__
    title = _rich_text(getattr(chart, "title", None))
    series_payload: list[dict[str, object]] = []
    for series_index, series in enumerate(getattr(chart, "series", []), start=1):
        categories_reference = _first_nested(
            series,
            (("cat",), ("xVal",)),
        )
        values_reference = _first_nested(
            series,
            (("val",), ("yVal",)),
        )
        category_formula, categories = _reference_values(
            workbook, categories_reference
        )
        value_formula, values = _reference_values(workbook, values_reference)
        size_formula, sizes = _reference_values(
            workbook, getattr(series, "zVal", None)
        )
        series_payload.append(
            {
                "name": _series_name(workbook, series, series_index),
                "category_formula": category_formula,
                "categories": categories,
                "value_formula": value_formula,
                "values": values,
                "size_formula": size_formula,
                "sizes": sizes,
            }
        )

    anchor = getattr(chart, "anchor", None)
    anchor_text = str(anchor) if isinstance(anchor, str) else None
    return {
        "sheet": sheet.title,
        "index": index,
        "type": chart_type,
        "title": title or f"{chart_type} {index}",
        "anchor": anchor_text,
        "bar_direction": getattr(chart, "type", None),
        "series": series_payload,
    }


def _chart_text(payload: dict[str, object]) -> str:
    lines = [
        "CHART_DATA",
        f"sheet: {_quoted(payload['sheet'])}",
        f"chart_index: {payload['index']}",
        f"chart_type: {payload['type']}",
        f"title: {_quoted(payload['title'])}",
    ]
    if payload.get("anchor"):
        lines.append(f"anchor: {payload['anchor']}")
    for index, series in enumerate(payload["series"], start=1):
        lines.extend(
            [
                f"series_{index}_name: {_quoted(series['name'])}",
                f"series_{index}_category_formula: {_quoted(series['category_formula'])}",
                "series_{}_categories: {}".format(
                    index,
                    json.dumps(series["categories"], ensure_ascii=False, default=str),
                ),
                f"series_{index}_value_formula: {_quoted(series['value_formula'])}",
                "series_{}_values: {}".format(
                    index,
                    json.dumps(series["values"], ensure_ascii=False, default=str),
                ),
            ]
        )
        if series["sizes"]:
            lines.append(
                "series_{}_sizes: {}".format(
                    index,
                    json.dumps(series["sizes"], ensure_ascii=False, default=str),
                )
            )
    return "\n".join(lines)


def _numbers(values: list[object]) -> list[float]:
    converted: list[float] = []
    for value in values:
        try:
            converted.append(float(value))
        except (TypeError, ValueError):
            converted.append(float("nan"))
    return converted


def _render_chart(payload: dict[str, object], output_path: Path) -> None:
    import matplotlib

    matplotlib.use("Agg")
    from matplotlib import pyplot as plt

    chart_type = str(payload["type"]).lower()
    series_items = payload["series"]
    figure = plt.figure(figsize=(8, 5), dpi=120)
    if "radar" in chart_type:
        axis = figure.add_subplot(111, polar=True)
    else:
        axis = figure.add_subplot(111)

    plotted = False
    for series_index, series in enumerate(series_items):
        values = _numbers(series["values"])
        if not values:
            continue
        labels = [str(value) for value in series["categories"]]
        if len(labels) != len(values):
            labels = [str(index + 1) for index in range(len(values))]
        positions = list(range(len(values)))
        name = str(series["name"])

        if "pie" in chart_type or "doughnut" in chart_type:
            safe_values = [0.0 if math.isnan(value) else value for value in values]
            wedges, *_ = axis.pie(safe_values, labels=labels, autopct="%1.1f%%")
            if "doughnut" in chart_type:
                centre = plt.Circle((0, 0), 0.55, color="white")
                axis.add_artist(centre)
            plotted = bool(wedges)
            break
        if "bar" in chart_type and payload.get("bar_direction") == "bar":
            axis.barh(positions, values, label=name)
            axis.set_yticks(positions, labels)
        elif "bar" in chart_type:
            width = 0.8 / max(1, len(series_items))
            offsets = [position + series_index * width for position in positions]
            axis.bar(offsets, values, width=width, label=name)
            axis.set_xticks(
                [position + width * (len(series_items) - 1) / 2 for position in positions],
                labels,
            )
        elif "scatter" in chart_type or "bubble" in chart_type:
            x_values = _numbers(series["categories"])
            if len(x_values) != len(values):
                x_values = [float(position) for position in positions]
            sizes = _numbers(series["sizes"])
            sizes = [max(10.0, value) for value in sizes] if len(sizes) == len(values) else 35
            axis.scatter(x_values, values, s=sizes, label=name)
        elif "area" in chart_type:
            axis.fill_between(positions, values, alpha=0.35, label=name)
            axis.plot(positions, values)
            axis.set_xticks(positions, labels)
        elif "radar" in chart_type:
            angles = [2 * math.pi * index / len(values) for index in positions]
            closed_values = values + values[:1]
            closed_angles = angles + angles[:1]
            axis.plot(closed_angles, closed_values, label=name)
            axis.fill(closed_angles, closed_values, alpha=0.15)
            axis.set_xticks(angles, labels)
        else:
            axis.plot(positions, values, marker="o", label=name)
            axis.set_xticks(positions, labels)
        plotted = True

    if not plotted:
        axis.text(0.5, 0.5, "Chart data unavailable", ha="center", va="center")
        axis.set_axis_off()
    elif "pie" not in chart_type and "doughnut" not in chart_type:
        axis.grid(True, alpha=0.25)
        if len(series_items) > 1:
            axis.legend()
    axis.set_title(str(payload["title"]))
    figure.tight_layout()
    figure.savefig(output_path, format="png", transparent=True)
    plt.close(figure)


def _extract_spreadsheet(source: Path, output_directory: Path) -> dict[str, object]:
    from openpyxl import load_workbook

    keep_vba = source.suffix.lower() in {".xlsm", ".xltm"}
    workbook = load_workbook(source, data_only=False, keep_vba=keep_vba)
    cached_workbook = load_workbook(source, data_only=True, keep_vba=keep_vba)
    text_blocks: list[dict[str, object]] = []
    visuals: list[dict[str, object]] = []
    issues: list[dict[str, str]] = []
    order = 0

    for sheet in workbook.worksheets:
        cached_sheet = cached_workbook[sheet.title]
        table_text = _worksheet_table_text(sheet, cached_sheet)
        if table_text is not None:
            text_blocks.append(
                {
                    "kind": "table",
                    "text": table_text,
                    "source_order": order,
                    "source_path": f"sheet:{sheet.title}",
                }
            )
            order += 1

        for chart_index, chart in enumerate(sheet._charts, start=1):
            payload = _chart_payload(cached_workbook, sheet, chart, chart_index)
            text_blocks.append(
                {
                    "kind": "chart_data",
                    "text": _chart_text(payload),
                    "source_order": order,
                    "source_path": f"sheet:{sheet.title}/chart:{chart_index}",
                }
            )
            safe_sheet = _SAFE_NAME.sub("_", sheet.title).strip("_") or "sheet"
            rendered = output_directory / f"{safe_sheet}_chart_{chart_index}.png"
            try:
                _render_chart(payload, rendered)
                visuals.append(
                    {
                        "kind": "rendered_chart",
                        "path": str(rendered.resolve()),
                        "source_order": order,
                        "source_path": f"sheet:{sheet.title}/chart:{chart_index}",
                    }
                )
            except Exception as error:
                issues.append(_issue("chart_render", f"{sheet.title} chart {chart_index}: {error}"))
            order += 1

    workbook.close()
    cached_workbook.close()
    return {
        "format": source.suffix[1:].upper(),
        "text_blocks": text_blocks,
        "visuals": visuals,
        "issues": issues,
    }


def _svg_bytes(source: Path) -> bytes:
    data = source.read_bytes()
    if source.suffix.lower() == ".svgz":
        data = gzip.decompress(data)
    upper = data[:65536].upper()
    if b"<!DOCTYPE" in upper or b"<!ENTITY" in upper:
        raise ValueError("SVG files containing DTD or entity declarations are rejected.")
    return data


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _svg_text(data: bytes) -> tuple[str, ET.Element]:
    root = ET.fromstring(data)
    lines = ["VECTOR_DATA", "format: SVG"]
    for name in ("width", "height", "viewBox", "preserveAspectRatio"):
        if name in root.attrib:
            lines.append(f"{name}: {_quoted(root.attrib[name])}")
    title_index = 0
    element_index = 0
    vector_names = {
        "path", "rect", "circle", "ellipse", "line", "polyline", "polygon",
        "text", "use", "image", "g", "symbol", "clipPath", "mask",
        "linearGradient", "radialGradient", "pattern", "marker",
    }
    for element in root.iter():
        name = _local_name(element.tag)
        if name in {"title", "desc"}:
            text = " ".join(part.strip() for part in element.itertext() if part.strip())
            if text:
                title_index += 1
                lines.append(f"{name}_{title_index}: {_quoted(text)}")
        if name not in vector_names:
            continue
        element_index += 1
        attributes = " ".join(
            f"{_local_name(key)}={_quoted(value)}"
            for key, value in sorted(element.attrib.items())
        )
        visible_text = " ".join(
            part.strip() for part in element.itertext() if part.strip()
        ) if name == "text" else ""
        suffix = f" text={_quoted(visible_text)}" if visible_text else ""
        lines.append(f"element_{element_index}: {name} {attributes}{suffix}".rstrip())
    return "\n".join(lines), root


def _sanitized_svg(root: ET.Element) -> tuple[bytes, bool]:
    safe_root = ET.fromstring(ET.tostring(root, encoding="utf-8"))
    changed = False

    def clean(parent: ET.Element) -> None:
        nonlocal changed
        for child in list(parent):
            if _local_name(child.tag) in {"script", "foreignObject"}:
                parent.remove(child)
                changed = True
                continue
            clean(child)
        for key, value in list(parent.attrib.items()):
            name = _local_name(key).lower()
            lower_value = value.strip().lower()
            if name.startswith("on"):
                del parent.attrib[key]
                changed = True
            elif name == "href" and not (
                lower_value.startswith("#") or lower_value.startswith("data:image/")
            ):
                del parent.attrib[key]
                changed = True
            elif name == "style" and "url(" in lower_value and "url(#" not in lower_value:
                del parent.attrib[key]
                changed = True
        if _local_name(parent.tag) == "style" and parent.text:
            cleaned = re.sub(r"@import\s+[^;]+;?", "", parent.text, flags=re.IGNORECASE)
            cleaned = re.sub(
                r"url\((?!\s*['\"]?#)[^)]+\)",
                "none",
                cleaned,
                flags=re.IGNORECASE,
            )
            if cleaned != parent.text:
                parent.text = cleaned
                changed = True

    clean(safe_root)
    return ET.tostring(safe_root, encoding="utf-8", xml_declaration=True), changed


def _svg_dimension(value: str | None) -> float | None:
    if value is None:
        return None
    match = re.fullmatch(r"\s*([0-9]+(?:\.[0-9]+)?)\s*(px|pt|pc|in|cm|mm)?\s*", value)
    if not match:
        return None
    number = float(match.group(1))
    unit = match.group(2) or "px"
    factors = {"px": 1.0, "pt": 96 / 72, "pc": 16.0, "in": 96.0, "cm": 96 / 2.54, "mm": 96 / 25.4}
    return number * factors[unit]


def _svg_size(root: ET.Element) -> tuple[int, int]:
    width = _svg_dimension(root.attrib.get("width"))
    height = _svg_dimension(root.attrib.get("height"))
    view_box = root.attrib.get("viewBox", "").replace(",", " ").split()
    if (width is None or height is None) and len(view_box) == 4:
        try:
            width = width or float(view_box[2])
            height = height or float(view_box[3])
        except ValueError:
            pass
    width = width or 1024
    height = height or 768
    return max(1, min(4096, round(width))), max(1, min(4096, round(height)))


def _resolve_browser(explicit: str | None) -> str:
    candidates = [
        explicit,
        os.environ.get("SVG_RENDERER_EXECUTABLE"),
        shutil.which("msedge"),
        shutil.which("chrome"),
        shutil.which("google-chrome"),
        shutil.which("chromium"),
        shutil.which("chromium-browser"),
    ]
    if os.name == "nt":
        candidates.extend(
            [
                str(Path(os.environ.get("PROGRAMFILES", "")) / "Google/Chrome/Application/chrome.exe"),
                str(Path(os.environ.get("PROGRAMFILES(X86)", "")) / "Microsoft/Edge/Application/msedge.exe"),
                str(Path(os.environ.get("PROGRAMFILES", "")) / "Microsoft/Edge/Application/msedge.exe"),
            ]
        )
    else:
        candidates.extend(
            [
                "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
                "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
            ]
        )
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return str(Path(candidate))
    raise FileNotFoundError(
        "A Chromium-compatible SVG renderer was not found. Supply "
        "--browser-executable or SVG_RENDERER_EXECUTABLE."
    )


def _render_svg(data: bytes, root: ET.Element, output_path: Path, browser: str | None) -> None:
    executable = _resolve_browser(browser)
    width, height = _svg_size(root)
    with tempfile.TemporaryDirectory(prefix="structured-svg-") as temporary:
        temporary_path = Path(temporary)
        source = temporary_path / "source.svg"
        profile = temporary_path / "browser-profile"
        source.write_bytes(data)
        command = [
            executable,
            "--headless=new",
            "--disable-gpu",
            "--disable-extensions",
            "--disable-background-networking",
            "--no-first-run",
            "--no-default-browser-check",
            "--hide-scrollbars",
            "--default-background-color=00000000",
            "--host-resolver-rules=MAP * 0.0.0.0",
            f"--user-data-dir={profile}",
            f"--window-size={width},{height}",
            f"--screenshot={output_path.resolve()}",
            source.resolve().as_uri(),
        ]
        completed = subprocess.run(
            command,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=60,
        )
        if completed.returncode != 0 or not output_path.is_file():
            detail = (completed.stderr or completed.stdout).strip()
            raise RuntimeError(detail or f"SVG renderer exited with {completed.returncode}.")


def _extract_vector(source: Path, output_directory: Path, browser: str | None) -> dict[str, object]:
    data = _svg_bytes(source)
    text, root = _svg_text(data)
    render_data, sanitized = _sanitized_svg(root)
    rendered = output_directory / "vector_render.png"
    issues: list[dict[str, str]] = []
    if sanitized:
        issues.append(
            _issue(
                "vector_safety",
                "Executable or external SVG content was excluded from rendering.",
                "warning",
            )
        )
    visuals: list[dict[str, object]] = []
    try:
        _render_svg(render_data, root, rendered, browser)
        visuals.append(
            {
                "kind": "rendered_vector",
                "path": str(rendered.resolve()),
                "source_order": 1,
                "source_path": "svg:rendered",
            }
        )
    except Exception as error:
        issues.append(_issue("vector_render", str(error)))
    return {
        "format": "SVGZ" if source.suffix.lower() == ".svgz" else "SVG",
        "text_blocks": [
            {
                "kind": "vector_data",
                "text": text,
                "source_order": 0,
                "source_path": "svg:source",
            }
        ],
        "visuals": visuals,
        "issues": issues,
    }


def extract(source_path: str, output_path: str, browser: str | None = None) -> dict[str, object]:
    source = Path(source_path)
    if not source.is_file():
        raise FileNotFoundError(f"Source file not found: {source}")
    output_directory = Path(output_path)
    output_directory.mkdir(parents=True, exist_ok=True)
    extension = source.suffix.lower()
    if extension in SPREADSHEET_EXTENSIONS:
        return _extract_spreadsheet(source, output_directory)
    if extension in VECTOR_EXTENSIONS:
        return _extract_vector(source, output_directory, browser)
    raise ValueError(f"Unsupported structured visual file: {source}")


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--browser-executable")
    return parser


def main() -> int:
    arguments = _parser().parse_args()
    try:
        payload = extract(
            arguments.input,
            arguments.output_dir,
            arguments.browser_executable,
        )
    except Exception as error:
        print(json.dumps({"error": str(error), "error_type": type(error).__name__}))
        return 1
    print(json.dumps(payload, ensure_ascii=True, separators=(",", ":"), default=str))
    return 0


if __name__ == "__main__":
    sys.exit(main())
