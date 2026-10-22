from __future__ import annotations

import base64
from email.message import EmailMessage
import importlib.util
from pathlib import Path
import tempfile
import unittest
from unittest import mock
import zipfile

import pymupdf as fitz
from PIL import Image
from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[3]
SERVICE_PATH = (
    ROOT / "ReleaseCode" / "services" / "compound_document" / "compound_document_reader.py"
)
SPEC = importlib.util.spec_from_file_location("compound_document_reader", SERVICE_PATH)
if SPEC is None or SPEC.loader is None:
    raise RuntimeError(f"Cannot load compound-document service: {SERVICE_PATH}")
READER = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(READER)


class CompoundDocumentReaderTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.directory = Path(self.temporary_directory.name)
        self.output = self.directory / "output"

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    def test_tika_jar_is_discovered_next_to_the_tika_command(self) -> None:
        command = self.directory / "tika.bat"
        command.write_text("@echo off", encoding="utf-8")
        older = self.directory / "tika-app-3.2.0.jar"
        current = self.directory / "tika-app-4.0.0.jar"
        older.write_bytes(b"older")
        current.write_bytes(b"current")

        with mock.patch.dict("os.environ", {"TIKA_APP_JAR": ""}), mock.patch.object(
            READER.shutil,
            "which",
            return_value=str(command),
        ):
            resolved = READER._resolve_tika_jar(None)

        self.assertEqual(resolved, str(current))

    def test_html_extracts_text_and_local_and_data_images(self) -> None:
        local_image = self.directory / "local.png"
        Image.new("RGB", (12, 12), "red").save(local_image)
        data_image = base64.b64encode(local_image.read_bytes()).decode("ascii")
        svg = base64.b64encode(
            b"<svg xmlns='http://www.w3.org/2000/svg' width='10' height='10'>"
            b"<path d='M0 0 L10 10' stroke='black'/></svg>"
        ).decode("ascii")
        source = self.directory / "sample.html"
        source.write_text(
            f"<html><body><p>direct text</p><img src='local.png'>"
            f"<img src='data:image/png;base64,{data_image}'>"
            f"<img src='data:image/svg+xml;base64,{svg}'></body></html>",
            encoding="utf-8",
        )

        result = READER.extract_document(str(source), str(self.output))

        self.assertEqual(result["detected_format"], "HTML")
        self.assertIn("direct text", result["text_blocks"][0]["text"])
        self.assertEqual(len(result["image_parts"]), 3)
        for image in result["image_parts"]:
            self.assertTrue(Path(image["file_path"]).is_file())
        self.assertEqual(result["audio_parts"], [])
        self.assertTrue(
            any(block["kind"] == "vector_data" for block in result["text_blocks"])
        )
        self.assertTrue(
            any(part["kind"] == "rendered_vector" for part in result["image_parts"])
        )

    def test_pdf_extracts_direct_text_embedded_image_and_page_pixels(self) -> None:
        image = self.directory / "image.png"
        Image.new("RGB", (20, 20), "blue").save(image)
        source = self.directory / "sample.pdf"
        document = fitz.open()
        page = document.new_page(width=200, height=200)
        page.insert_text((20, 30), "PDF direct text")
        page.draw_line((20, 40), (180, 40), color=(1, 0, 0), width=2)
        page.insert_image(fitz.Rect(20, 50, 80, 110), filename=str(image))
        document.save(source)
        document.close()

        result = READER.extract_document(str(source), str(self.output))

        self.assertEqual(result["detected_format"], "PDF")
        self.assertTrue(
            any("PDF direct text" in block["text"] for block in result["text_blocks"])
        )
        self.assertTrue(
            any(part["kind"] == "rendered_page" for part in result["image_parts"])
        )
        self.assertTrue(
            any(part["kind"] == "embedded_image" for part in result["image_parts"])
        )
        self.assertTrue(
            any(block["kind"] == "vector_data" for block in result["text_blocks"])
        )

    def test_pptx_preserves_slide_text_and_image_relationship(self) -> None:
        image = self.directory / "slide-image.png"
        Image.new("RGB", (20, 20), "green").save(image)
        source = self.directory / "sample.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        text_box.text = "slide direct text"
        slide.shapes.add_picture(str(image), Inches(1), Inches(2))
        presentation.save(source)

        result = READER.extract_document(str(source), str(self.output))

        self.assertEqual(result["detected_format"], "PPTX")
        self.assertTrue(
            any(
                "slide direct text" in block["text"] and block["page_index"] == 0
                for block in result["text_blocks"]
            )
        )
        self.assertTrue(
            any(
                part["page_index"] == 0
                and str(part["logical_location"]).startswith("slide:1/")
                for part in result["image_parts"]
            )
        )

    def test_ods_keeps_table_and_chart_source_and_renders_chart_pixels(self) -> None:
        source = self.directory / "chart.ods"
        content = """<?xml version='1.0' encoding='UTF-8'?>
<office:document-content
 xmlns:office='urn:oasis:names:tc:opendocument:xmlns:office:1.0'
 xmlns:table='urn:oasis:names:tc:opendocument:xmlns:table:1.0'
 xmlns:text='urn:oasis:names:tc:opendocument:xmlns:text:1.0'
 xmlns:chart='urn:oasis:names:tc:opendocument:xmlns:chart:1.0'>
 <office:body><office:spreadsheet>
  <table:table table:name='Data'>
   <table:table-row>
    <table:table-cell office:value-type='string'><text:p>Month</text:p></table:table-cell>
    <table:table-cell office:value-type='float' office:value='1'><text:p>1</text:p></table:table-cell>
   </table:table-row>
   <table:table-row>
    <table:table-cell office:value-type='string'><text:p>Value</text:p></table:table-cell>
    <table:table-cell office:value-type='float' office:value='4'><text:p>4</text:p></table:table-cell>
   </table:table-row>
  </table:table>
  <chart:chart chart:class='chart:line'/>
 </office:spreadsheet></office:body>
</office:document-content>"""
        with zipfile.ZipFile(source, "w") as archive:
            archive.writestr(
                "mimetype",
                "application/vnd.oasis.opendocument.spreadsheet",
            )
            archive.writestr("content.xml", content)

        result = READER.extract_document(str(source), str(self.output))

        self.assertEqual(result["detected_format"], "ODS")
        self.assertTrue(
            any(block["kind"] == "table" for block in result["text_blocks"])
        )
        self.assertTrue(
            any(block["kind"] == "chart_data" for block in result["text_blocks"])
        )
        self.assertTrue(
            any(part["kind"] == "rendered_chart" for part in result["image_parts"])
        )

    def test_fb2_zip_extracts_text_and_embedded_image_without_archive_recursion(self) -> None:
        image = self.directory / "cover.png"
        Image.new("RGB", (12, 12), "purple").save(image)
        encoded = base64.b64encode(image.read_bytes()).decode("ascii")
        xml = (
            "<FictionBook xmlns='http://www.gribuser.ru/xml/fictionbook/2.0'>"
            "<body><section><p>book text</p></section></body>"
            f"<binary id='cover.png' content-type='image/png'>{encoded}</binary>"
            "</FictionBook>"
        )
        source = self.directory / "book.fb2.zip"
        with zipfile.ZipFile(source, "w") as archive:
            archive.writestr("book.fb2", xml)
            archive.writestr("ignored.zip", b"PK\x03\x04")

        result = READER.extract_document(str(source), str(self.output))

        self.assertEqual(result["detected_format"], "FB2_ZIP")
        self.assertTrue(any("book text" in item["text"] for item in result["text_blocks"]))
        self.assertEqual(len(result["image_parts"]), 1)
        self.assertEqual(result["audio_parts"], [])

    def test_format_geometry_relationship_and_origin_helpers(self) -> None:
        self.assertEqual(READER._format_label(Path("book.fb2.zip")), "FB2_ZIP")
        self.assertEqual(READER._format_label(Path("unknown.xyz")), "XYZ")

        class Rectangle:
            x0, y0, x1, y1 = 1, 2, 3, 4

        class Point:
            x, y = 5, 6

        self.assertEqual(READER._json_safe_vector(None), None)
        self.assertEqual(READER._json_safe_vector({"p": Point()}), {"p": {"x": 5.0, "y": 6.0}})
        self.assertEqual(
            READER._json_safe_vector((Rectangle(), "value")),
            [{"x0": 1.0, "y0": 2.0, "x1": 3.0, "y1": 4.0}, "value"],
        )
        self.assertEqual(READER._json_safe_vector(object()).startswith("<object object"), True)

        self.assertEqual(READER._relationship_source("_rels/.rels"), "")
        self.assertEqual(
            READER._relationship_source("ppt/slides/_rels/slide1.xml.rels"),
            "ppt/slides/slide1.xml",
        )
        self.assertIsNone(READER._relationship_source("plain.xml"))
        reverse = {
            "ppt/media/a.png": ["ppt/slides/slide2.xml"],
            "xl/media/b.png": ["xl/worksheets/sheet3.xml"],
            "word/media/c.png": ["word/document.xml"],
            "media/deep.png": ["links/one.xml"],
            "links/one.xml": ["ppt/slides/slide4.xml"],
        }
        self.assertEqual(READER._media_origins("ppt/media/a.png", reverse)[0][0], 1)
        self.assertIn("sheet:3", READER._media_origins("xl/media/b.png", reverse)[0][1])
        self.assertIn("word/document.xml", READER._media_origins("word/media/c.png", reverse)[0][1])
        self.assertEqual(READER._media_origins("missing.png", reverse), [(None, "missing.png")])
        self.assertEqual(READER._media_origins("media/deep.png", reverse)[0][0], 3)

    def test_relationship_and_reference_helpers_reject_unsafe_inputs(self) -> None:
        archive_path = self.directory / "relationships.zip"
        relationships = """<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'>
          <Relationship Id='a' Target='../media/a.png'/>
          <Relationship Id='b' Target='https://example.invalid/a' TargetMode='External'/>
          <Relationship Id='c'/>
        </Relationships>"""
        with zipfile.ZipFile(archive_path, "w") as archive:
            archive.writestr("ppt/slides/_rels/slide1.xml.rels", relationships)
            archive.writestr("bad/_rels/item.xml.rels", b"not xml")
        with zipfile.ZipFile(archive_path) as archive:
            reverse = READER._package_relationships(archive)
        self.assertEqual(reverse["ppt/media/a.png"], ["ppt/slides/slide1.xml"])

        source = self.directory / "page.md"
        source.write_text("page", encoding="utf-8")
        local = self.directory / "asset.bin"
        local.write_bytes(b"asset")
        self.assertEqual(READER._safe_local_reference(source, "asset.bin"), local.resolve())
        self.assertIsNone(READER._safe_local_reference(source, "https://example.invalid/a"))
        self.assertIsNone(READER._safe_local_reference(source, "../outside.bin"))
        self.assertIsNone(READER._safe_local_reference(source, "missing.bin"))

        extraction = READER.Extraction(source, self.output)
        READER._asset_reference(extraction, "data:audio/wav;base64,YXVkaW8=", "inline")
        READER._asset_reference(extraction, "data:text/plain,hello", "ignored")
        READER._asset_reference(extraction, "data:image/png;base64", "bad")
        READER._asset_reference(extraction, "asset.bin", "local")
        self.assertEqual(len(extraction.audio_parts), 1)
        self.assertTrue(any(issue["stage"] == "data_uri" for issue in extraction.issues))

    def test_message_rtf_and_fb2_edge_paths(self) -> None:
        message = EmailMessage()
        message["Subject"] = "mixed"
        message.set_content("plain body")
        message.add_alternative("<p>html body</p>", subtype="html")
        message.add_attachment(b"image", maintype="image", subtype="png", filename="a.png")
        message.add_attachment(b"audio", maintype="audio", subtype="wav", filename="a.wav")
        source = self.directory / "message.eml"
        source.write_bytes(message.as_bytes())
        extraction = READER.Extraction(source, self.output / "message")
        READER._extract_message(extraction)
        self.assertTrue(any("plain body" in item["text"] for item in extraction.text_blocks))
        self.assertTrue(extraction.image_parts)
        self.assertTrue(extraction.audio_parts)

        rtf = self.directory / "sample.rtf"
        rtf.write_text(r"{\rtf1\ansi RTF text}", encoding="latin-1")
        rtf_extraction = READER.Extraction(rtf, self.output / "rtf")
        READER._extract_rtf(rtf_extraction)
        self.assertIn("RTF text", rtf_extraction.text_blocks[0]["text"])

        fb2 = self.directory / "edge.fb2"
        fb2_extraction = READER.Extraction(fb2, self.output / "fb2")
        with self.assertRaises(ValueError):
            READER._extract_fb2_xml(fb2_extraction, b"<!DOCTYPE x><FictionBook/>", "bad.fb2")
        with self.assertRaises(ValueError):
            READER._extract_fb2_xml(fb2_extraction, b"<Other/>", "bad.fb2")
        data = b"""<FictionBook><body><p>story</p></body>
          <binary id='skip.bin' content-type='application/octet-stream'>YQ==</binary>
          <binary id='bad.png' content-type='image/png'>!!!</binary>
          <binary id='sound' content-type='audio/wav'>YXVkaW8=</binary>
        </FictionBook>"""
        READER._extract_fb2_xml(fb2_extraction, data, "edge.fb2")
        self.assertEqual(len(fb2_extraction.audio_parts), 1)
        self.assertGreaterEqual(len(fb2_extraction.issues), 2)

    def test_dependency_and_tika_execution_branches(self) -> None:
        java_home = self.directory / "java"
        java = java_home / "bin" / ("java.exe" if READER.os.name == "nt" else "java")
        java.parent.mkdir(parents=True)
        java.write_bytes(b"")
        jar_home = self.directory / "tika"
        jar_home.mkdir()
        older = jar_home / "tika-app-1.jar"
        newer = jar_home / "tika-app-2.jar"
        older.write_bytes(b"")
        newer.write_bytes(b"")
        with mock.patch.dict(READER.os.environ, {"JAVA_HOME": str(java_home), "TIKA_HOME": str(jar_home)}, clear=True), mock.patch.object(READER.shutil, "which", return_value=None):
            self.assertEqual(READER._resolve_java(None), str(java))
            self.assertEqual(READER._resolve_tika_jar(None), str(newer))
        with mock.patch.dict(READER.os.environ, {}, clear=True), mock.patch.object(
            READER.shutil, "which", return_value=None
        ):
            self.assertIsNone(READER._resolve_java(str(self.directory / "missing")))

        source = self.directory / "tika.bin"
        source.write_bytes(b"source")
        extraction = READER.Extraction(source, self.output / "tika")
        READER._tika_text(extraction, None, None)
        self.assertTrue(extraction.issues)

        completed = mock.Mock(returncode=0, stdout="<html><body><table><tr><th>A</th><td>1</td></tr></table><p>body</p></body></html>", stderr="")
        with mock.patch.object(READER.subprocess, "run", return_value=completed):
            READER._tika_text(extraction, str(java), str(newer))
        self.assertTrue(any(item["kind"] == "table" for item in extraction.text_blocks))
        self.assertTrue(any("body" in item["text"] for item in extraction.text_blocks))

        failed = mock.Mock(returncode=2, stdout="", stderr="failed")
        with mock.patch.object(READER.subprocess, "run", return_value=failed):
            READER._tika_text(extraction, str(java), str(newer))
            READER._tika_embedded(extraction, str(java), str(newer))
        with mock.patch.object(READER.subprocess, "run", side_effect=OSError("broken")):
            READER._tika_text(extraction, str(java), str(newer))
            READER._tika_embedded(extraction, str(java), str(newer))
        self.assertTrue(any(item["stage"] == "tika_embedded" for item in extraction.issues))

    def test_markdown_package_chart_and_required_dependency_failures(self) -> None:
        image = self.directory / "markdown.png"
        Image.new("RGB", (4, 4), "orange").save(image)
        audio = self.directory / "sound.wav"
        audio.write_bytes(b"audio")
        markdown = self.directory / "sample.md"
        markdown.write_text(
            "# title\n![image](markdown.png)\n<audio src='sound.wav'>",
            encoding="utf-8",
        )
        extraction = READER.Extraction(markdown, self.output / "markdown")
        READER._extract_markdown(extraction)
        self.assertEqual(len(extraction.image_parts), 1)
        self.assertEqual(len(extraction.audio_parts), 1)

        package_path = self.directory / "chart.docx"
        chart_xml = b"""<c:chartSpace xmlns:c='urn:chart'><c:v>name</c:v><c:v>3.5</c:v><c:v>bad</c:v></c:chartSpace>"""
        with zipfile.ZipFile(package_path, "w") as archive:
            archive.writestr("word/charts/chart1.xml", chart_xml)
        package_extraction = READER.Extraction(package_path, self.output / "package")
        READER._extract_package(package_extraction)
        self.assertTrue(any(item["kind"] == "chart_data" for item in package_extraction.text_blocks))
        self.assertTrue(any(item["kind"] == "rendered_chart" for item in package_extraction.image_parts))

        broken = self.directory / "broken.docx"
        broken.write_bytes(b"not a zip")
        broken_extraction = READER.Extraction(broken, self.output / "broken")
        READER._extract_package(broken_extraction)
        self.assertEqual(broken_extraction.issues[0]["stage"], "package")

        with self.assertRaises(FileNotFoundError):
            READER.extract_document(str(self.directory / "missing.docx"), str(self.output))
        plain = self.directory / "plain.html"
        plain.write_text("<p>x</p>", encoding="utf-8")
        with mock.patch.object(READER, "_resolve_java", return_value=None), self.assertRaisesRegex(RuntimeError, "Java"):
            READER.extract_document(str(plain), str(self.output / "no-java"))
        with mock.patch.object(READER, "_resolve_java", return_value="java"), mock.patch.object(
            READER, "_resolve_tika_jar", return_value=None
        ), self.assertRaisesRegex(RuntimeError, "Tika"):
            READER.extract_document(str(plain), str(self.output / "no-tika"))

    def test_tika_embedded_exports_supported_files_and_can_skip_images(self) -> None:
        source = self.directory / "embedded.bin"
        source.write_bytes(b"source")
        extraction = READER.Extraction(source, self.output / "embedded")

        def create_embedded(command, **kwargs):
            argument = next(item for item in command if item.startswith("--extract-dir="))
            directory = Path(argument.split("=", 1)[1])
            (directory / "image.png").write_bytes(b"image")
            (directory / "audio.mp3").write_bytes(b"audio")
            return SimpleNamespace(returncode=0, stdout="", stderr="")

        from types import SimpleNamespace

        with mock.patch.object(READER.subprocess, "run", side_effect=create_embedded):
            READER._tika_embedded(extraction, "java", "tika.jar", include_images=False)
        self.assertEqual(len(extraction.image_parts), 0)
        self.assertEqual(len(extraction.audio_parts), 1)


if __name__ == "__main__":
    unittest.main()
