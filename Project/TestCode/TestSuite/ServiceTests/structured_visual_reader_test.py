from __future__ import annotations

import importlib.util
from pathlib import Path
import gzip
import tempfile
import unittest
from unittest import mock
from types import SimpleNamespace
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.table import Table
from PIL import Image


ROOT = Path(__file__).resolve().parents[3]
SERVICE_PATH = (
    ROOT / "ReleaseCode" / "services" / "structured_visual" / "structured_visual_reader.py"
)
SPEC = importlib.util.spec_from_file_location("structured_visual_reader", SERVICE_PATH)
if SPEC is None or SPEC.loader is None:
    raise RuntimeError(f"Cannot load structured-visual service: {SERVICE_PATH}")
READER = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(READER)


class StructuredVisualReaderTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.directory = Path(self.temporary_directory.name)

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    def _workbook(self) -> Path:
        path = self.directory / "chart.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sales"
        sheet.append(["Month", "Value"])
        sheet.append(["Jan", 10])
        sheet.append(["Feb", 20])
        sheet.append(["Mar", 15])
        sheet.add_table(Table(displayName="SalesData", ref="A1:B4"))
        chart = BarChart()
        chart.title = "Quarter sales"
        chart.add_data(
            Reference(sheet, min_col=2, min_row=1, max_row=4),
            titles_from_data=True,
        )
        chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=4))
        sheet.add_chart(chart, "D2")
        workbook.save(path)
        return path

    def test_spreadsheet_outputs_table_chart_text_and_rendered_chart(self) -> None:
        result = READER.extract(
            str(self._workbook()),
            str(self.directory / "output"),
        )

        self.assertEqual(result["format"], "XLSX")
        self.assertEqual(
            [block["kind"] for block in result["text_blocks"]],
            ["table", "chart_data"],
        )
        self.assertIn('A1="Month"', result["text_blocks"][0]["text"])
        self.assertIn("named_tables: SalesData=A1:B4", result["text_blocks"][0]["text"])
        self.assertIn("series_1_values: [10, 20, 15]", result["text_blocks"][1]["text"])
        self.assertEqual(result["visuals"][0]["kind"], "rendered_chart")
        rendered = Path(result["visuals"][0]["path"])
        self.assertTrue(rendered.is_file())
        with Image.open(rendered) as image:
            self.assertEqual(image.format, "PNG")
            self.assertGreater(image.width, 0)

    def test_svg_outputs_vector_text_and_rendered_pixels(self) -> None:
        source = self.directory / "sample.svg"
        source.write_text(
            '<svg xmlns="http://www.w3.org/2000/svg" width="120" height="80">'
            '<title>Vector sample</title><rect x="5" y="5" width="100" '
            'height="50" fill="red"/><text x="10" y="30">Hello</text></svg>',
            encoding="utf-8",
        )

        def fake_render(data, root, output_path, browser):
            Image.new("RGBA", (120, 80), "red").save(output_path)

        with mock.patch.object(READER, "_render_svg", side_effect=fake_render):
            result = READER.extract(str(source), str(self.directory / "svg-output"))

        self.assertEqual(result["format"], "SVG")
        self.assertEqual(result["text_blocks"][0]["kind"], "vector_data")
        self.assertIn("Vector sample", result["text_blocks"][0]["text"])
        self.assertIn("element_1: rect", result["text_blocks"][0]["text"])
        self.assertEqual(result["visuals"][0]["kind"], "rendered_vector")
        self.assertTrue(Path(result["visuals"][0]["path"]).is_file())

    def test_svg_rejects_entity_declarations(self) -> None:
        source = self.directory / "unsafe.svg"
        source.write_text(
            '<!DOCTYPE svg [<!ENTITY x "unsafe">]>'
            '<svg xmlns="http://www.w3.org/2000/svg"><text>&x;</text></svg>',
            encoding="utf-8",
        )

        with self.assertRaisesRegex(ValueError, "entity"):
            READER.extract(str(source), str(self.directory / "unsafe-output"))

    def test_svgz_is_decompressed_and_unsafe_content_is_not_rendered(self) -> None:
        source = self.directory / "sample.svgz"
        source.write_bytes(
            gzip.compress(
                b'<svg xmlns="http://www.w3.org/2000/svg" width="20" height="10">'
                b'<script>alert(1)</script><image href="file:///secret.png"/>'
                b'<rect width="20" height="10"/></svg>'
            )
        )
        captured: list[bytes] = []

        def fake_render(data, root, output_path, browser):
            captured.append(data)
            Image.new("RGBA", (20, 10), "blue").save(output_path)

        with mock.patch.object(READER, "_render_svg", side_effect=fake_render):
            result = READER.extract(str(source), str(self.directory / "svgz-output"))

        self.assertEqual(result["format"], "SVGZ")
        self.assertNotIn(b"script", captured[0].lower())
        self.assertNotIn(b"file:///secret.png", captured[0])
        self.assertEqual(result["issues"][0]["stage"], "vector_safety")

    def test_structured_helper_fallbacks_cover_empty_and_literal_data(self) -> None:
        self.assertEqual(READER._cell_text(None), "")
        self.assertEqual(READER._cell_text(True), "TRUE")
        self.assertEqual(READER._cell_text(False), "FALSE")
        self.assertEqual(READER._cell_text("a\r\nb\rc"), "a\nb\nc")
        self.assertEqual(READER._issue("x", "y", "warning")["severity"], "warning")
        self.assertIsNone(READER._nested(None, "missing"))
        self.assertIsNone(READER._first_nested(SimpleNamespace(), (("x",),)))
        self.assertIsNone(READER._rich_text(SimpleNamespace()))
        self.assertEqual(READER._range_values({}, None), [])
        self.assertEqual(READER._range_values({}, "missing!A1"), [])

        points = [
            SimpleNamespace(idx="2", v="c"),
            SimpleNamespace(idx="0", v="a"),
            SimpleNamespace(idx="1", v="b"),
        ]
        literal = SimpleNamespace(numLit=SimpleNamespace(pt=points))
        self.assertEqual(READER._literal_cache(literal, "numLit"), ["a", "b", "c"])
        self.assertEqual(READER._literal_cache(literal, "strLit"), [])
        formula, values = READER._reference_values({}, literal)
        self.assertIsNone(formula)
        self.assertEqual(values, ["a", "b", "c"])
        self.assertEqual(READER._reference_values({}, None), (None, []))
        self.assertEqual(
            READER._series_name({}, SimpleNamespace(tx=SimpleNamespace(v="Named")), 1),
            "Named",
        )
        self.assertEqual(READER._series_name({}, SimpleNamespace(), 4), "Series 4")

        empty = Workbook()
        self.assertIsNone(READER._worksheet_table_text(empty.active, None))

    def test_svg_safety_dimensions_browser_and_renderer_failures(self) -> None:
        data = b"""<svg xmlns='http://www.w3.org/2000/svg' width='2in' height='bad' viewBox='0 0 300 150'>
          <style>@import 'bad.css'; .a { fill:url(http://bad); }</style>
          <g onclick='bad()' style='fill:url(http://bad)'>
            <image href='https://bad.invalid/a.png'/><use href='#safe'/>
            <foreignObject/><script/></g><text>Hello</text></svg>"""
        text, root = READER._svg_text(data)
        cleaned, changed = READER._sanitized_svg(root)
        self.assertTrue(changed)
        self.assertIn("Hello", text)
        self.assertNotIn(b"script", cleaned)
        self.assertNotIn(b"https://bad.invalid", cleaned)
        self.assertEqual(READER._svg_dimension(None), None)
        self.assertEqual(READER._svg_dimension("bad"), None)
        self.assertEqual(READER._svg_dimension("2in"), 192.0)
        self.assertEqual(READER._svg_size(root), (192, 150))
        fallback = ET.fromstring("<svg viewBox='0 0 bad bad'/>")
        self.assertEqual(READER._svg_size(fallback), (1024, 768))

        browser = self.directory / "browser.exe"
        browser.write_bytes(b"")
        self.assertEqual(READER._resolve_browser(str(browser)), str(browser))
        with mock.patch.object(READER.shutil, "which", return_value=None), mock.patch.dict(
            READER.os.environ, {}, clear=True
        ), self.assertRaises(FileNotFoundError):
            READER._resolve_browser(str(self.directory / "missing"))

        output = self.directory / "shot.png"
        failed = SimpleNamespace(returncode=2, stderr="renderer failed", stdout="")
        with mock.patch.object(READER, "_resolve_browser", return_value=str(browser)), mock.patch.object(
            READER.subprocess, "run", return_value=failed
        ), self.assertRaisesRegex(RuntimeError, "renderer failed"):
            READER._render_svg(cleaned, root, output, str(browser))

    def test_chart_renderer_covers_supported_chart_families(self) -> None:
        common = {
            "title": "Chart",
            "bar_direction": None,
            "series": [
                {
                    "name": "one",
                    "categories": ["A", "B", "C"],
                    "values": [1, 2, 3],
                    "sizes": [5, 20, 30],
                },
                {
                    "name": "two",
                    "categories": [1, 2],
                    "values": [2, 4, 6],
                    "sizes": [],
                },
            ],
        }
        self.assertTrue(READER.math.isnan(READER._numbers(["bad"])[0]))
        for chart_type in ("PieChart", "DoughnutChart", "BarChart", "ScatterChart", "AreaChart", "RadarChart", "LineChart"):
            payload = dict(common, type=chart_type)
            if chart_type == "BarChart":
                payload["bar_direction"] = "bar"
            output = self.directory / f"{chart_type}.png"
            READER._render_chart(payload, output)
            self.assertTrue(output.is_file())

        empty = {
            "type": "LineChart",
            "title": "Empty",
            "bar_direction": None,
            "series": [{"name": "none", "categories": [], "values": [], "sizes": []}],
        }
        output = self.directory / "empty.png"
        READER._render_chart(empty, output)
        self.assertTrue(output.is_file())


if __name__ == "__main__":
    unittest.main()
